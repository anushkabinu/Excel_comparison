"""
report_generator.py
-------------------
Generates a timestamped, colour-coded Excel comparison report with 5 sheets:
  1. Summary
  2. Column Mapping
  3. Row Comparison
  4. Cell Differences
  5. Missing & Extra Rows
"""

from __future__ import annotations

import datetime
import os
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from reconciler.comparator import ComparisonResult


# ---------------------------------------------------------------------------
# Colour palette
# ---------------------------------------------------------------------------

_GREEN  = PatternFill("solid", fgColor="C6EFCE")   # matched
_RED    = PatternFill("solid", fgColor="FFC7CE")   # mismatch / extra
_YELLOW = PatternFill("solid", fgColor="FFEB9C")   # missing / warning
_BLUE   = PatternFill("solid", fgColor="BDD7EE")   # header
_GREY   = PatternFill("solid", fgColor="D9D9D9")   # sub-header
_WHITE  = PatternFill("solid", fgColor="FFFFFF")
_ORANGE = PatternFill("solid", fgColor="FFD966")   # low confidence

_HEADER_FONT = Font(bold=True, color="1F3864")
_TITLE_FONT  = Font(bold=True, size=14, color="1F3864")

_THIN_BORDER_SIDE = Side(style="thin", color="BFBFBF")
_THIN_BORDER = Border(
    left=_THIN_BORDER_SIDE, right=_THIN_BORDER_SIDE,
    top=_THIN_BORDER_SIDE,  bottom=_THIN_BORDER_SIDE,
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _write_header_row(ws, row_idx: int, values: list, fill=_BLUE):
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = fill
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER


def _write_data_row(ws, row_idx: int, values: list, fill=None):
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        if fill:
            cell.fill = fill
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        cell.border = _THIN_BORDER


def _autofit_columns(ws, min_width: int = 10, max_width: int = 50):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


def _freeze_top_row(ws):
    ws.freeze_panes = ws.cell(row=2, column=1)


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _build_summary_sheet(ws, result: ComparisonResult):
    s = result.summary
    ws.title = "Summary"

    # Title
    ws.merge_cells("A1:D1")
    title_cell = ws["A1"]
    title_cell.value = "Excel Reconciliation — Summary"
    title_cell.font = _TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    title_cell.fill = _BLUE
    ws.row_dimensions[1].height = 28

    metrics = [
        ("", "", "", ""),
        ("Row Statistics", "", "", ""),
        ("Client rows", s.total_client_rows, "Output rows", s.total_output_rows),
        ("Matched rows", s.matched_rows, "Mismatched rows", s.mismatched_rows),
        ("Missing rows (client not in output)", s.missing_rows,
         "Extra rows (output not in client)", s.extra_rows),
        ("", "", "", ""),
        ("Cell Statistics", "", "", ""),
        ("Total cells compared", s.total_cells_compared,
         "Cell match %", f"{s.match_pct:.2f}%"),
        ("Matched cells", s.matched_cells,
         "Mismatched cells", s.mismatched_cells),
        ("Missing cell values", s.missing_cell_values, "", ""),
        ("", "", "", ""),
        ("Tolerance Settings", "", "", ""),
        ("Absolute tolerance", s.abs_tolerance, "Relative tolerance", s.rel_tolerance),
    ]

    sub_header_rows = {3, 8, 13}

    for row_offset, row_data in enumerate(metrics, start=2):
        row_idx = row_offset
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = _THIN_BORDER
            if row_offset in sub_header_rows and col_idx == 1:
                cell.font = Font(bold=True, color="1F3864")
                cell.fill = _GREY
            cell.alignment = Alignment(vertical="center")

    _autofit_columns(ws)


def _build_column_mapping_sheet(ws, result: ComparisonResult):
    ws.title = "Column Mapping"
    headers = [
        "Client Column", "Output Column", "Confidence %",
        "Name Score", "Synonym Score", "Type Score",
        "Overlap Score", "Uniqueness Score", "Status",
    ]
    _write_header_row(ws, 1, headers)
    _freeze_top_row(ws)

    all_mappings = result.column_mappings
    for row_idx, m in enumerate(all_mappings, start=2):
        conf_pct = f"{m.confidence * 100:.1f}%"
        status = getattr(m, "status", "auto")
        if m.output_col is None:
            fill = _YELLOW
            status = "unmapped"
        elif m.confidence >= 0.80:
            fill = _GREEN
        elif m.confidence >= 0.50:
            fill = _ORANGE
        else:
            fill = _RED

        _write_data_row(ws, row_idx, [
            m.client_col,
            m.output_col or "— unmapped —",
            conf_pct,
            f"{m.name_score:.2f}",
            f"{m.synonym_score:.2f}",
            f"{m.type_score:.2f}",
            f"{m.overlap_score:.2f}",
            f"{m.unique_score:.2f}",
            status.upper(),
        ], fill=fill)

    _autofit_columns(ws)


def _build_row_comparison_sheet(ws, result: ComparisonResult):
    ws.title = "Row Comparison"
    headers = [
        "Row Key", "Client Excel Row", "Output Excel Row",
        "Status", "Mismatched Columns",
    ]
    _write_header_row(ws, 1, headers)
    _freeze_top_row(ws)

    status_fill = {
        "MATCH":    _GREEN,
        "MISMATCH": _RED,
        "MISSING":  _YELLOW,
        "EXTRA":    _ORANGE,
    }

    for row_idx, rr in enumerate(result.row_results, start=2):
        fill = status_fill.get(rr.status, _WHITE)
        _write_data_row(ws, row_idx, [
            rr.row_key,
            rr.client_excel_row or "",
            rr.output_excel_row or "",
            rr.status,
            ", ".join(rr.mismatch_columns) if rr.mismatch_columns else "",
        ], fill=fill)

    _autofit_columns(ws)


def _build_cell_differences_sheet(ws, result: ComparisonResult):
    ws.title = "Cell Differences"
    headers = [
        "Row Key",
        "Client Excel Row", "Output Excel Row",
        "Column",
        "Client Value", "Output Value",
        "Numeric Diff", "Diff %",
        "Issue",
    ]
    _write_header_row(ws, 1, headers)
    _freeze_top_row(ws)

    issue_fill = {
        "MISMATCH":       _RED,
        "MISSING_CLIENT": _YELLOW,
        "MISSING_OUTPUT": _YELLOW,
    }

    for row_idx, cd in enumerate(result.cell_diffs, start=2):
        fill = issue_fill.get(cd.issue, _RED)
        ndiff = f"{cd.numeric_diff:.6g}" if cd.numeric_diff is not None else ""
        dpct  = f"{cd.diff_pct:.4f}%" if cd.diff_pct is not None else ""
        _write_data_row(ws, row_idx, [
            cd.row_key,
            cd.client_excel_row or "",
            cd.output_excel_row or "",
            cd.column_name,
            cd.client_value,
            cd.output_value,
            ndiff,
            dpct,
            cd.issue,
        ], fill=fill)

    _autofit_columns(ws)


def _build_missing_extra_sheet(ws, result: ComparisonResult):
    ws.title = "Missing & Extra Rows"

    current_row = 1

    def _write_section(title: str, df: pd.DataFrame, fill):
        nonlocal current_row
        # Section title
        ws.cell(row=current_row, column=1, value=title).font = Font(bold=True, size=12)
        ws.cell(row=current_row, column=1).fill = _GREY
        current_row += 1

        if df.empty:
            ws.cell(row=current_row, column=1, value="(none)").font = Font(italic=True)
            current_row += 2
            return

        # Drop internal columns for display
        display_df = df.drop(
            columns=[c for c in df.columns if c.startswith("_")],
            errors="ignore",
        )

        # Header
        _write_header_row(ws, current_row, display_df.columns.tolist(), fill=_BLUE)
        current_row += 1

        for _, row in display_df.iterrows():
            _write_data_row(ws, current_row, row.tolist(), fill=fill)
            current_row += 1

        current_row += 1  # blank spacer row

    _write_section(
        f"Missing Rows ({len(result.missing_rows_df)}) — In Client, Not in Output",
        result.missing_rows_df, _YELLOW,
    )
    _write_section(
        f"Extra Rows ({len(result.extra_rows_df)}) — In Output, Not in Client",
        result.extra_rows_df, _ORANGE,
    )

    _autofit_columns(ws)


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate_report(
    result: ComparisonResult,
    output_dir: str = ".",
    base_name: str = "reconciliation_report",
) -> str:
    """
    Write a timestamped Excel report and return the file path.

    Parameters
    ----------
    result      : ComparisonResult from comparator.compare_dataframes()
    output_dir  : directory to write the report into
    base_name   : prefix for the filename

    Returns
    -------
    Absolute path to the generated .xlsx file
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{base_name}_{timestamp}.xlsx"
    filepath = os.path.join(output_dir, filename)

    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Build each sheet
    _build_summary_sheet(wb.create_sheet(), result)
    _build_column_mapping_sheet(wb.create_sheet(), result)
    _build_row_comparison_sheet(wb.create_sheet(), result)
    _build_cell_differences_sheet(wb.create_sheet(), result)
    _build_missing_extra_sheet(wb.create_sheet(), result)

    wb.save(filepath)
    return os.path.abspath(filepath)


# ---------------------------------------------------------------------------
# Wide-format report  (ACTUAL | EXPECTED | RESULT per column)
# ---------------------------------------------------------------------------

def _safe_cell_value(idx_df: pd.DataFrame, excel_row, col_name: str):
    """Look up a single cell by excel row number; return '' on any miss."""
    if excel_row is None or col_name not in idx_df.columns:
        return ""
    try:
        val = idx_df.at[excel_row, col_name]
    except Exception:
        return ""
    # .at can return a Series if the index has duplicates — take first
    if isinstance(val, pd.Series):
        val = val.iloc[0] if len(val) else ""
    if pd.isna(val):
        return ""
    return val


def generate_wide_report(
    result: ComparisonResult,
    df_client: pd.DataFrame,
    df_output: pd.DataFrame,
    output_dir: str = ".",
    base_name: str = "wide_comparison_report",
) -> str:
    """
    Generate a wide-format Excel report where every compared column appears
    as three sub-columns: ACTUAL (client), EXPECTED (output), RESULT (PASS/FAIL).

    Each data row is a **key-matched pair** from ComparisonResult (not positional).
    ACTUAL is always taken from the client excel row; EXPECTED from the paired
    output excel row. PASS/FAIL comes from the comparator (incl. numeric tolerance).

    Layout
    ------
    Row 1 : ROW #  |  OUT ROW #  |  <FIELD_NAME>  (merged 3 cols)  |  ...
    Row 2 :        |             |  ACTUAL  |  EXPECTED  |  RESULT  |  ...
    Row 3+: data
    Footer: per-field match rates + average match rate

    Sheet 2 — Field Match Summary: total / matched / not matched / match rate per field.

    Returns the absolute path to the written .xlsx file.
    """
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    filename  = f"{base_name}_{timestamp}.xlsx"
    filepath  = os.path.join(output_dir, filename)

    # ── 1. Pre-build lookup structures ─────────────────────────────────────

    # Index original DataFrames by _excel_row_ for fast value lookup
    _ERR = "_excel_row_"
    if _ERR in df_client.columns:
        client_idx = df_client.set_index(_ERR)
    else:
        client_idx = df_client.copy()
        client_idx.index = range(len(client_idx))

    if _ERR in df_output.columns:
        output_idx = df_output.set_index(_ERR)
    else:
        output_idx = df_output.copy()
        output_idx.index = range(len(output_idx))

    # Cells the comparator flagged as different (already applied numeric tolerance)
    diff_set: set = {(cd.row_key, cd.column_name) for cd in result.cell_diffs}

    # Only columns that have an output mapping (excluded cols are already filtered out)
    col_pairs = [
        (m.client_col, m.output_col)
        for m in result.column_mappings
        if m.output_col
    ]

    # ── 2. Build data rows from key-matched pairs only ─────────────────────

    data_rows = []
    field_pass_counts = [0] * len(col_pairs)

    for rr in result.row_results:
        if rr.status in ("MISSING", "EXTRA"):
            continue

        # Always pull ACTUAL/EXPECTED from the exact key-paired excel rows
        row = [rr.client_excel_row, rr.output_excel_row]

        for field_idx, (client_col, output_col) in enumerate(col_pairs):
            c_val = _safe_cell_value(client_idx, rr.client_excel_row, client_col)
            o_val = _safe_cell_value(output_idx, rr.output_excel_row, output_col)

            # RESULT follows comparator decision (tolerance already applied)
            if (rr.row_key, client_col) in diff_set:
                res = "FAIL"
            else:
                res = "PASS"
                field_pass_counts[field_idx] += 1

            row.extend([c_val, o_val, res])

        data_rows.append(row)

    n_compared = len(data_rows)
    field_match_rates = [
        (field_pass_counts[i] / n_compared * 100.0) if n_compared > 0 else 0.0
        for i in range(len(col_pairs))
    ]
    avg_match_rate = (
        sum(field_match_rates) / len(field_match_rates)
        if field_match_rates else 0.0
    )

    # ── 3. Write workbook ──────────────────────────────────────────────────

    wb = Workbook()
    ws = wb.active
    ws.title = "Side-by-Side"

    # Colour fills
    _PASS_FILL   = PatternFill("solid", fgColor="C6EFCE")   # green
    _FAIL_FILL   = PatternFill("solid", fgColor="FFC7CE")   # red
    _FIELD_FILL  = PatternFill("solid", fgColor="BDD7EE")   # blue — field name header
    _SUB_FILL    = PatternFill("solid", fgColor="D9D9D9")   # grey — ACTUAL/EXPECTED/RESULT
    _FOOTER_FILL = PatternFill("solid", fgColor="FFF2CC")   # light yellow — match rates

    bold_blue = Font(bold=True, color="1F3864")
    bold_dark = Font(bold=True, color="1F3864")
    center    = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Column layout: A = Client ROW #, B = Output ROW #, then field triples from col 3
    _DATA_COL0 = 3   # first ACTUAL column

    # ── Row 1: field name headers (merged across 3 columns each) ───────────
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    c = ws.cell(row=1, column=1, value="ROW #")
    c.font = bold_blue
    c.fill = _FIELD_FILL
    c.alignment = center
    c.border = _THIN_BORDER

    ws.merge_cells(start_row=1, start_column=2, end_row=2, end_column=2)
    c = ws.cell(row=1, column=2, value="OUT ROW #")
    c.font = bold_blue
    c.fill = _FIELD_FILL
    c.alignment = center
    c.border = _THIN_BORDER

    for field_idx, (client_col, _output_col) in enumerate(col_pairs):
        col_start = _DATA_COL0 + field_idx * 3
        col_end   = col_start + 2

        ws.merge_cells(start_row=1, start_column=col_start, end_row=1, end_column=col_end)
        c = ws.cell(row=1, column=col_start, value=client_col)
        c.font = bold_blue
        c.fill = _FIELD_FILL
        c.alignment = center
        c.border = _THIN_BORDER

    # ── Row 2: ACTUAL / EXPECTED / RESULT sub-headers ─────────────────────
    for field_idx in range(len(col_pairs)):
        col_start = _DATA_COL0 + field_idx * 3
        for sub_idx, label in enumerate(["ACTUAL", "EXPECTED", "RESULT"]):
            c = ws.cell(row=2, column=col_start + sub_idx, value=label)
            c.font = bold_dark
            c.fill = _SUB_FILL
            c.alignment = center
            c.border = _THIN_BORDER

    # ── Rows 3+: data ─────────────────────────────────────────────────────
    for row_offset, row_vals in enumerate(data_rows, start=3):
        for col_offset, val in enumerate(row_vals, start=1):
            c = ws.cell(row=row_offset, column=col_offset, value=val)
            c.border = _THIN_BORDER
            if col_offset <= 2:
                c.alignment = center
                continue
            # Colour the RESULT column (third of each triple)
            sub_pos = (col_offset - _DATA_COL0) % 3
            if sub_pos == 2:
                c.fill = _PASS_FILL if val == "PASS" else _FAIL_FILL
                c.font = Font(bold=True)
                c.alignment = center

    last_data_row = 2 + n_compared   # header rows 1-2, data starts at 3

    # ── Footer: per-field match rate + average ─────────────────────────────
    footer_row = last_data_row + 2
    c = ws.cell(row=footer_row, column=1, value="MATCH RATE")
    c.font = Font(bold=True, size=11, color="1F3864")
    c.fill = _FOOTER_FILL
    c.alignment = center
    c.border = _THIN_BORDER
    ws.merge_cells(
        start_row=footer_row, start_column=1,
        end_row=footer_row, end_column=2,
    )

    for field_idx, rate in enumerate(field_match_rates):
        col_start = _DATA_COL0 + field_idx * 3
        col_end   = col_start + 2
        ws.merge_cells(
            start_row=footer_row, start_column=col_start,
            end_row=footer_row, end_column=col_end,
        )
        rate_fill = (
            _PASS_FILL if rate >= 95 else
            _FOOTER_FILL if rate >= 80 else
            _FAIL_FILL
        )
        c = ws.cell(row=footer_row, column=col_start, value=f"{rate:.1f}%")
        c.font = Font(bold=True)
        c.fill = rate_fill
        c.alignment = center
        c.border = _THIN_BORDER

    avg_row = footer_row + 1
    c = ws.cell(
        row=avg_row, column=1,
        value=(
            f"Average Match Rate: {avg_match_rate:.1f}%   "
            f"({n_compared} paired rows x {len(col_pairs)} fields)   "
            f"| abs_tol={result.summary.abs_tolerance}  "
            f"rel_tol={result.summary.rel_tolerance}"
        ),
    )
    c.font = Font(bold=True, size=12, color="1F3864")
    c.fill = _FOOTER_FILL
    end_col = max(2, _DATA_COL0 + len(col_pairs) * 3 - 1)
    ws.merge_cells(
        start_row=avg_row, start_column=1,
        end_row=avg_row, end_column=end_col,
    )

    # ── Sheet 2: field match summary ───────────────────────────────────────
    ws_sum = wb.create_sheet("Field Match Summary")
    sum_headers = [
        "Field Name",
        "Total Fields",
        "PASS",
        "FAIL",
        "ACCURACY",
        "AVERAGE ACCURACY"
    ]
    _write_header_row(ws_sum, 1, sum_headers)
    _freeze_top_row(ws_sum)

    n_fields = len(col_pairs)
    data_start = 2                     # first data row (row 2)
    data_end   = data_start + n_fields - 1  # last data row

    for field_idx, (client_col, _) in enumerate(col_pairs):
        matched     = field_pass_counts[field_idx]
        not_matched = n_compared - matched
        rate        = field_match_rates[field_idx]
        _write_data_row(ws_sum, data_start + field_idx, [
            client_col,
            n_compared,
            matched,
            not_matched,
            f"{rate:.2f}%",
            "",   # column F left blank — filled by merged cell below
        ])

    # ── Single merged AVERAGE ACCURACY cell spanning all data rows ──────────
    if n_fields > 0:
        _AVG_FILL  = PatternFill("solid", fgColor="FFF2CC")   # light yellow
        _avg_color = "375623" if avg_match_rate >= 95 else "9C5700" if avg_match_rate >= 80 else "9C0006"

        if data_end > data_start:
            ws_sum.merge_cells(
                start_row=data_start, start_column=6,
                end_row=data_end,     end_column=6,
            )
        c = ws_sum.cell(row=data_start, column=6,
                        value=f"{avg_match_rate:.2f}%")
        c.font      = Font(bold=True, size=14, color=_avg_color)
        c.fill      = _AVG_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = _THIN_BORDER

    _autofit_columns(ws_sum)
    ws_sum.column_dimensions["F"].width = 20




    # ── Column widths ──────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 10   # Client ROW #
    ws.column_dimensions["B"].width = 12   # Output ROW #
    for field_idx in range(len(col_pairs)):
        col_start = _DATA_COL0 + field_idx * 3
        for sub in range(3):
            letter = get_column_letter(col_start + sub)
            ws.column_dimensions[letter].width = 30 if sub < 2 else 10

    # Freeze first two header rows + row-number columns
    ws.freeze_panes = ws.cell(row=3, column=_DATA_COL0)

    wb.save(filepath)
    return os.path.abspath(filepath)

